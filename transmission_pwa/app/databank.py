from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import gspread
from google.oauth2.service_account import Credentials
from openpyxl import load_workbook

from .config import get_settings
from .models import DatabankSource


APP_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = APP_ROOT / "data"
SOURCES_FILE = DATA_DIR / "sources.json"
LOCAL_WORKBOOK = PROJECT_ROOT / "5. OSOGBO RCC - (330_132kV LINES) LOAD MANAGEMENT TRACKING - MAY 2026.xlsx"
DEFAULT_SHEET_ID = "1jsxcy4UcJtn5cCXFC--zBrcexRD8oSvoelHiI9_aXRE"
DEFAULT_SOURCE_ID = "osogbo-rcc-330-132kv-load-management-may-2026"
READONLY_SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]


@dataclass
class CacheEntry:
    loaded_at: float
    records: list[dict[str, Any]]


_CACHE: dict[str, CacheEntry] = {}


def ensure_default_sources() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if SOURCES_FILE.exists():
        return
    sources = [
        DatabankSource(
            id=DEFAULT_SOURCE_ID,
            name="Osogbo RCC 330/132kV Lines Load Management Tracking",
            description=(
                "Daily hourly load-management tracker for Osogbo Regional Control Centre "
                "330kV and 132kV transmission lines. Contains ACC, transmission interface, "
                "line voltage, line nomenclature, DISCO, monthly maximum references, and "
                "01:00-24:00 MW or operational-status readings."
            ),
            sheet_id=DEFAULT_SHEET_ID,
            sheet_url=f"https://docs.google.com/spreadsheets/d/{DEFAULT_SHEET_ID}/edit",
            category="Transmission line load management",
            voltage_levels=["330kV", "132kV"],
            tags=[
                "Osogbo RCC",
                "load management",
                "330kV lines",
                "132kV lines",
                "hourly readings",
                "DISCO",
            ],
        ).model_dump(mode="json")
    ]
    SOURCES_FILE.write_text(json.dumps(sources, indent=2), encoding="utf-8")


def list_sources() -> list[DatabankSource]:
    ensure_default_sources()
    raw = json.loads(SOURCES_FILE.read_text(encoding="utf-8"))
    return [DatabankSource(**item) for item in raw]


def save_source(source: DatabankSource) -> DatabankSource:
    sources = list_sources()
    if any(item.id == source.id for item in sources):
        raise ValueError("A source with this id already exists.")
    sources.append(source)
    SOURCES_FILE.write_text(
        json.dumps([item.model_dump(mode="json") for item in sources], indent=2),
        encoding="utf-8",
    )
    return source


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "source"


def extract_sheet_id(value: str) -> str:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", value)
    if match:
        return match.group(1)
    cleaned = value.strip()
    if cleaned and "/" not in cleaned and " " not in cleaned:
        return cleaned
    raise ValueError("Provide a valid Google Sheet URL or raw Sheet ID.")


def load_records_for_sources(source_ids: list[str] | None = None) -> list[dict[str, Any]]:
    settings = get_settings()
    selected = [source for source in list_sources() if source.active]
    if source_ids:
        requested = set(source_ids)
        selected = [source for source in selected if source.id in requested]

    records: list[dict[str, Any]] = []
    for source in selected:
        cached = _CACHE.get(source.id)
        if cached and time.time() - cached.loaded_at < settings.data_cache_seconds:
            records.extend(cached.records)
            continue
        loaded = load_source_records(source)
        _CACHE[source.id] = CacheEntry(time.time(), loaded)
        records.extend(loaded)
    return records


def load_source_records(source: DatabankSource) -> list[dict[str, Any]]:
    settings = get_settings()
    if source.id == DEFAULT_SOURCE_ID and settings.prefer_local_workbook and LOCAL_WORKBOOK.exists():
        return parse_workbook_records(LOCAL_WORKBOOK, source)
    try:
        return parse_google_sheet_records(source)
    except Exception:
        if source.id == DEFAULT_SOURCE_ID and LOCAL_WORKBOOK.exists():
            return parse_workbook_records(LOCAL_WORKBOOK, source)
        raise


def parse_google_sheet_records(source: DatabankSource) -> list[dict[str, Any]]:
    settings = get_settings()
    credentials = Credentials.from_service_account_file(
        settings.google_service_account_file,
        scopes=READONLY_SCOPES,
    )
    client = gspread.authorize(credentials)
    spreadsheet = client.open_by_key(source.sheet_id)
    records: list[dict[str, Any]] = []
    for worksheet in spreadsheet.worksheets():
        values = worksheet.get_all_values()
        records.extend(parse_grid_records(values, worksheet.title, source))
    return records


def parse_workbook_records(path: Path, source: DatabankSource) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    records: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = [
            [cell for cell in row]
            for row in worksheet.iter_rows(min_row=1, max_row=80, min_col=1, max_col=35, values_only=True)
        ]
        records.extend(parse_grid_records(rows, worksheet.title, source))
    return records


def parse_grid_records(rows: list[list[Any]], sheet_name: str, source: DatabankSource) -> list[dict[str, Any]]:
    if not rows or not ordinal_day(sheet_name):
        return []
    header_index = find_header_row(rows)
    if header_index is None:
        return []
    header = [clean_text(value) for value in rows[header_index]]
    hour_columns = [(idx, normalize_hour(label)) for idx, label in enumerate(header) if normalize_hour(label)]
    if not hour_columns:
        return []

    sheet_date = extract_sheet_date(rows, sheet_name)
    records: list[dict[str, Any]] = []
    last_acc = ""
    last_interface = ""
    last_voltage = ""

    for row in rows[header_index + 1 :]:
        serial = cell(row, 0)
        if to_float(clean_text(serial)) is None:
            continue
        line_name = clean_text(cell(row, 4))
        if not line_name:
            continue

        last_acc = clean_text(cell(row, 1)) or last_acc
        last_interface = clean_text(cell(row, 2)) or last_interface
        last_voltage = clean_text(cell(row, 3)) or last_voltage
        if last_voltage not in {"330kV", "132kV"}:
            continue

        base = {
            "source_id": source.id,
            "source_name": source.name,
            "sheet_name": sheet_name,
            "date": sheet_date.isoformat() if sheet_date else None,
            "serial_number": clean_text(serial),
            "acc": last_acc,
            "transmission_interface": last_interface,
            "line_voltage": last_voltage,
            "line_nomenclature": line_name,
            "disco": clean_text(cell(row, 10)),
            "max_for_month": clean_text(cell(row, 5)),
            "max_for_month_time_date": clean_text(cell(row, 6)),
            "new_monthly_max": clean_text(cell(row, 7)),
            "max_load_ever": clean_text(cell(row, 8)),
            "max_load_ever_time_date": clean_text(cell(row, 9)),
        }

        for idx, hour in hour_columns:
            raw_value = clean_text(cell(row, idx))
            if not raw_value:
                continue
            mw_value = to_float(raw_value)
            records.append(
                {
                    **base,
                    "hour": hour,
                    "load_mw": mw_value,
                    "status": None if mw_value is not None else raw_value,
                    "raw_value": raw_value,
                }
            )
    return records


def find_header_row(rows: list[list[Any]]) -> int | None:
    for idx, row in enumerate(rows[:10]):
        labels = {clean_text(value).lower() for value in row}
        if "s/n" in labels and "line nomenclature" in labels and "disco" in labels:
            return idx
    return None


def extract_sheet_date(rows: list[list[Any]], sheet_name: str) -> date | None:
    day = ordinal_day(sheet_name)
    if day:
        return date(2026, 5, day)
    for row in rows[:3]:
        for value in row:
            text = clean_text(value)
            match = re.search(r"(\d{2})/(\d{2})/(\d{4})", text)
            if match:
                day, month, year = map(int, match.groups())
                return date(year, month, day)
    return None


def ordinal_day(sheet_name: str) -> int | None:
    match = re.fullmatch(r"(\d{1,2})(?:ST|ND|RD|TH)", sheet_name.upper().strip())
    if not match:
        return None
    day = int(match.group(1))
    return day if 1 <= day <= 31 else None


def normalize_hour(value: str) -> str | None:
    text = clean_text(value)
    if re.fullmatch(r"\d{1,2}:\d{2}", text):
        hour = int(text.split(":", 1)[0])
        if 1 <= hour <= 24:
            return f"{hour:02d}:00"
    try:
        numeric = float(text)
    except ValueError:
        return None
    if 0 < numeric <= 1:
        total_minutes = round(numeric * 24 * 60)
        hour = total_minutes // 60
        minute = total_minutes % 60
        if minute == 0 and 1 <= hour <= 24:
            return f"{hour:02d}:00"
    return None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    text = str(value).replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def cell(row: list[Any], idx: int) -> Any:
    return row[idx] if idx < len(row) else None


def to_float(value: str) -> float | None:
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return None


def build_catalog(records: list[dict[str, Any]]) -> dict[str, Any]:
    numeric = [record for record in records if record["load_mw"] is not None]
    statuses = [record for record in records if record["status"]]
    return {
        "records": len(records),
        "numeric_readings": len(numeric),
        "status_readings": len(statuses),
        "dates": sorted({record["date"] for record in records if record["date"]}),
        "accs": sorted({record["acc"] for record in records if record["acc"]}),
        "interfaces": sorted({record["transmission_interface"] for record in records if record["transmission_interface"]}),
        "voltages": sorted({record["line_voltage"] for record in records if record["line_voltage"]}),
        "discos": sorted({record["disco"] for record in records if record["disco"]}),
        "lines": sorted({record["line_nomenclature"] for record in records if record["line_nomenclature"]}),
    }
